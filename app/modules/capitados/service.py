import uuid
import io
from datetime import date, datetime, timedelta
from typing import List, Optional, Dict, Any
from sqlalchemy import select, and_, or_
from sqlalchemy.ext.asyncio import AsyncSession
from openpyxl import load_workbook
from loguru import logger

from app.modules.capitados.models import (
    CapitadosBatchLog,
    CapitadosBatchItemLog,
    CapitadosProductInsured,
    CapitadosContract,
    CapitadosMonthlyRecord,
)
from app.modules.plans.models import (
    Product,
    Plan,
    PlanVersion,
    PlanVersionAgeSurcharge,
    PlanVersionCountry,
)
from app.modules.geography.models import Country


class CapitadosBatchService:
    @staticmethod
    async def create_batch_log(
        db: AsyncSession,
        company_id: uuid.UUID,
        coverage_month: date,
        original_filename: str,
        user_id: Optional[uuid.UUID] = None,
    ) -> CapitadosBatchLog:
        batch = CapitadosBatchLog(
            company_id=company_id,
            coverage_month=coverage_month,
            original_filename=original_filename,
            created_by_user_id=user_id,
            status="processing",
            reconciliation_status="pendiente",
        )
        db.add(batch)
        await db.commit()
        await db.refresh(batch)
        return batch

    @staticmethod
    async def process_batch(db: AsyncSession, batch_id: uuid.UUID, file_content: bytes):
        batch = await db.get(CapitadosBatchLog, batch_id)
        if not batch:
            logger.error(f"Batch {batch_id} not found")
            return

        try:
            wb = load_workbook(io.BytesIO(file_content), read_only=True, data_only=True)
            sheet = wb.active

            # Identify headers
            headers = {}
            for i, cell in enumerate(
                next(sheet.iter_rows(max_row=1, values_only=True))
            ):
                if cell:
                    headers[str(cell).strip().lower()] = i

            rows = list(sheet.iter_rows(min_row=2, values_only=True))
            batch.total_rows = len(rows)
            await db.commit()

            for idx, row in enumerate(rows):
                row_number = idx + 2
                await CapitadosBatchService._process_row(
                    db, batch, row, row_number, headers
                )

                # Commit every 50 rows to keep memory low and database healthy
                if idx % 50 == 0:
                    await db.commit()

            batch.status = "completed"
            batch.processed_at = datetime.now()

            # Simple reconciliation logic: if any rejected, status = con_diferencias
            if batch.total_rejected > 0:
                batch.reconciliation_status = "con_diferencias"
            else:
                batch.reconciliation_status = "conciliado"

            await db.commit()

        except Exception as e:
            logger.exception(f"Error processing batch {batch_id}")
            batch.status = "failed"
            batch.error_summary = str(e)
            await db.commit()

    @staticmethod
    async def _process_row(
        db: AsyncSession,
        batch: CapitadosBatchLog,
        row: tuple,
        row_number: int,
        headers: Dict[str, int],
    ):
        def get_val(key_aliases: List[str]) -> Any:
            for alias in key_aliases:
                if alias in headers:
                    return row[headers[alias]]
            return None

        # Extract data
        doc_number = str(
            get_val(["documento", "document_number", "document"]) or ""
        ).strip()
        full_name = str(get_val(["nombre", "full_name", "name"]) or "").strip()
        sex = str(get_val(["sexo", "sex", "gender"]) or "").strip().upper()[:1]
        residence_iso = (
            str(get_val(["pais_residencia", "residence_country", "residence"]) or "")
            .strip()
            .upper()
        )
        repatriation_iso = (
            str(
                get_val(["pais_repatriacion", "repatriation_country", "repatriation"])
                or ""
            )
            .strip()
            .upper()
        )
        age = get_val(["edad", "age"])
        product_id_raw = get_val(["producto_id", "product_id"])
        version_num = get_val(["version_id", "version", "version_number"])

        item_log = CapitadosBatchItemLog(
            batch_id=batch.id,
            row_number=row_number,
            document_number=doc_number,
            full_name=full_name,
            sex=sex,
            age_reported=age,
            residence_raw=residence_iso,
            repatriation_raw=repatriation_iso,
        )
        db.add(item_log)

        # Basic validations
        if not doc_number or not full_name or not product_id_raw:
            item_log.result = "rejected"
            item_log.rejection_code = "missing_required_fields"
            item_log.rejection_detail = "Document, Name, and Product ID are required."
            batch.total_rejected += 1
            return

        try:
            # 1. Resolve Product & Version
            try:
                prod_uuid = uuid.UUID(str(product_id_raw))
                prod_stmt = select(Product).where(Product.id == prod_uuid)
            except ValueError:
                prod_stmt = select(Product).where(Product.name == str(product_id_raw))

            product = (await db.execute(prod_stmt)).scalar_one_or_none()
            if not product:
                item_log.result = "rejected"
                item_log.rejection_code = "product_not_found"
                batch.total_rejected += 1
                return

            item_log.product_id = product.id

            # Resolve Version (latest if not provided)
            ver_stmt = (
                select(PlanVersion).join(Plan).where(Plan.product_id == product.id)
            )
            if version_num:
                ver_stmt = ver_stmt.where(
                    PlanVersion.version_number == int(version_num)
                )
            else:
                ver_stmt = ver_stmt.order_by(PlanVersion.version_number.desc())

            plan_version = (await db.execute(ver_stmt.limit(1))).scalar_one_or_none()
            if not plan_version:
                item_log.result = "rejected"
                item_log.rejection_code = "plan_version_not_found"
                batch.total_rejected += 1
                return

            item_log.plan_version_id = plan_version.id

            # 2. Resolve Countries
            res_country = (
                await db.execute(
                    select(Country).where(
                        or_(
                            Country.iso2 == residence_iso, Country.iso3 == residence_iso
                        )
                    )
                )
            ).scalar_one_or_none()
            rep_country = (
                await db.execute(
                    select(Country).where(
                        or_(
                            Country.iso2 == repatriation_iso,
                            Country.iso3 == repatriation_iso,
                        )
                    )
                )
            ).scalar_one_or_none()

            if res_country:
                item_log.residence_country_id = res_country.id
            if rep_country:
                item_log.repatriation_country_id = rep_country.id

            # 3. Find or Create Insured

            insured_stmt = select(CapitadosProductInsured).where(
                and_(
                    CapitadosProductInsured.company_id == batch.company_id,
                    CapitadosProductInsured.product_id == product.id,
                    CapitadosProductInsured.document_number == doc_number,
                    CapitadosProductInsured.rolled_back_at.is_(None),
                )
            )
            insured = (await db.execute(insured_stmt)).scalar_one_or_none()

            if not insured:
                insured = CapitadosProductInsured(
                    company_id=batch.company_id,
                    product_id=product.id,
                    document_number=doc_number,
                    full_name=full_name,
                    sex=sex,
                    residence_country_id=res_country.id if res_country else None,
                    repatriation_country_id=rep_country.id if rep_country else None,
                    age_reported=age,
                    status="active",
                )
                db.add(insured)
                await db.flush()  # Get ID

            item_log.person_id = insured.id

            # 4. Find or Create Contract
            contract_stmt = select(CapitadosContract).where(
                and_(
                    CapitadosContract.person_id == insured.id,
                    CapitadosContract.product_id == product.id,
                    CapitadosContract.rolled_back_at.is_(None),
                )
            )
            contract = (await db.execute(contract_stmt)).scalar_one_or_none()

            if not contract:
                # Calculate waiting periods
                vesting = CapitadosBatchService._calculate_vesting(
                    batch.coverage_month, plan_version
                )
                contract = CapitadosContract(
                    uuid=str(uuid.uuid4()),
                    company_id=batch.company_id,
                    product_id=product.id,
                    person_id=insured.id,
                    status="active",
                    entry_date=batch.coverage_month,
                    entry_age=age,
                    **vesting,
                )
                db.add(contract)
                await db.flush()

            item_log.contract_id = contract.id

            # 5. Check for Duplicate Monthly Record
            dup_stmt = select(CapitadosMonthlyRecord).where(
                and_(
                    CapitadosMonthlyRecord.person_id == insured.id,
                    CapitadosMonthlyRecord.product_id == product.id,
                    CapitadosMonthlyRecord.coverage_month == batch.coverage_month,
                    CapitadosMonthlyRecord.rolled_back_at.is_(None),
                )
            )
            existing_record = (await db.execute(dup_stmt)).scalar_one_or_none()

            if existing_record:
                item_log.result = "duplicated"
                item_log.duplicated_record_id = existing_record.id
                batch.total_duplicated += 1
                return

            # 6. Pricing
            pricing = await CapitadosBatchService._get_pricing(
                db, plan_version, residence_iso, age
            )

            # 7. Create Monthly Record
            monthly_record = CapitadosMonthlyRecord(
                company_id=batch.company_id,
                product_id=product.id,
                person_id=insured.id,
                contract_id=contract.id,
                coverage_month=batch.coverage_month,
                plan_version_id=plan_version.id,
                load_batch_id=batch.id,
                full_name=full_name,
                sex=sex,
                age_reported=age,
                residence_country_id=res_country.id if res_country else None,
                repatriation_country_id=rep_country.id if rep_country else None,
                **pricing,
                status="active",
            )
            db.add(monthly_record)
            await db.flush()

            item_log.monthly_record_id = monthly_record.id
            item_log.result = "applied"
            batch.total_applied += 1

        except Exception as e:
            logger.error(f"Error processing row {row_number}: {e}")
            item_log.result = "rejected"
            item_log.rejection_code = "internal_error"
            item_log.rejection_detail = str(e)
            batch.total_rejected += 1

    @staticmethod
    def _calculate_vesting(entry_date: date, version: PlanVersion) -> Dict[str, date]:
        return {
            "wtime_suicide_ends_at": entry_date + timedelta(days=version.wtime_suicide),
            "wtime_preexisting_conditions_ends_at": entry_date
            + timedelta(days=version.wtime_preexisting),
            "wtime_accident_ends_at": entry_date
            + timedelta(days=version.wtime_accident),
        }

    @staticmethod
    async def _get_pricing(
        db: AsyncSession, version: PlanVersion, country_iso: str, age: Optional[int]
    ) -> Dict[str, Any]:
        # Base Price from Country Overrides or Plan Version
        base_price = version.public_price or 0
        price_source = "plan_version"

        if country_iso:
            country_price = (
                await db.execute(
                    select(PlanVersionCountry).where(
                        and_(
                            PlanVersionCountry.plan_version_id == version.id,
                            PlanVersionCountry.country_code == country_iso,
                            PlanVersionCountry.is_available.is_(True),
                        )
                    )
                )
            ).scalar_one_or_none()

            if country_price and country_price.price_override:
                base_price = country_price.price_override
                price_source = "country_override"

        # Surcharge logic
        surcharge_pct = 0
        surcharge_rule_id = None
        surcharge_amount = 0

        if age is not None:
            surcharge = (
                await db.execute(
                    select(PlanVersionAgeSurcharge).where(
                        and_(
                            PlanVersionAgeSurcharge.plan_version_id == version.id,
                            PlanVersionAgeSurcharge.min_age <= age,
                            PlanVersionAgeSurcharge.max_age >= age,
                        )
                    )
                )
            ).scalar_one_or_none()

            if surcharge:
                surcharge_pct = surcharge.surcharge_percentage
                surcharge_rule_id = surcharge.id
                surcharge_amount = (float(base_price) * float(surcharge_pct)) / 100

        return {
            "price_base": base_price,
            "price_source": price_source,
            "age_surcharge_rule_id": surcharge_rule_id,
            "age_surcharge_percent": surcharge_pct,
            "age_surcharge_amount": surcharge_amount,
            "price_final": float(base_price) + float(surcharge_amount),
        }
